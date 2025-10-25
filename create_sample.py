#!/usr/bin/env python3
"""
Skripta za kreiranje primjer Excel fajla sa više sheet-ova i cenama.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create_sample_excel():
    """Kreira primjer Excel fajla sa više sheet-ova."""

    wb = Workbook()
    wb.remove(wb.active)  # Ukloni default sheet

    # Definišemo proizvode za različite sheet-ove
    sheets_data = {
        "Januar 2024": [
            ["Šifra", "Naziv proizvoda", "Cena (RSD)", "Količina", "Popust (%)"],
            ["001", "Laptop Dell XPS 13", 120000, 15, 5],
            ["002", "Monitor LG 27''", 35000, 30, 0],
            ["003", "Tastatura mehanička", 8500, 50, 10],
            ["004", "Miš Logitech MX", 6500, 40, 0],
            ["005", "USB-C Hub", 4500, 25, 15],
        ],
        "Februar 2024": [
            ["Šifra", "Naziv proizvoda", "Cena (RSD)", "Količina", "Popust (%)"],
            ["001", "Laptop Dell XPS 13", 120000, 12, 5],
            ["002", "Monitor LG 27''", 35000, 28, 0],
            ["003", "Tastatura mehanička", 8500, 45, 10],
            ["004", "Miš Logitech MX", 6500, 38, 0],
            ["005", "USB-C Hub", 4500, 22, 15],
            ["006", "WebCam HD", 7500, 20, 5],
        ],
        "Mart 2024": [
            ["Šifra", "Naziv", "Maloprodajna cena", "Veleprodajna cena", "Na stanju"],
            ["101", "Slušalice Sony WH-1000XM4", 35000, 28000, 18],
            ["102", "AirPods Pro", 28000, 22000, 25],
            ["103", "Jabra Elite 85t", 22000, 17000, 15],
            ["104", "Samsung Buds Pro", 18000, 14000, 30],
        ],
        "Specijalna ponuda": [
            ["Proizvod", "Redovna cena", "Akcijska cena", "Uštedjeno"],
            ["Gaming PC RTX 4070", 180000, 155000, 25000],
            ["MacBook Air M2", 165000, 145000, 20000],
            ["iPad Pro 12.9''", 125000, 110000, 15000],
            ["iPhone 15 Pro", 145000, 132000, 13000],
            ["Samsung S24 Ultra", 135000, 120000, 15000],
        ]
    }

    # Kreiraj sheet-ove
    for sheet_name, data in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)

        # Header stil
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center")

        # Dodaj podatke
        for row_idx, row_data in enumerate(data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)

                # Stilizuj header
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment

        # Podesi širinu kolona
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    # Sačuvaj fajl
    output_file = "cenovnik_primer.xlsx"
    wb.save(output_file)
    print(f"✅ Kreiran primjer Excel fajl: {output_file}")
    print(f"📊 Sheet-ovi: {', '.join(sheets_data.keys())}")
    print(f"\n💡 Sada možeš testirati:")
    print(f"   python update_prices.py {output_file} 1.1 --dry-run")


if __name__ == '__main__':
    create_sample_excel()
