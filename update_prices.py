#!/usr/bin/env python3
"""
Skripta za automatsko ažuriranje cena u Excel fajlovima.
Automatski prolazi kroz sve sheet-ove i ažurira cene prema zadatom koeficijentu.
"""

import argparse
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import re


def is_price_column(header):
    """Proverava da li naziv kolone ukazuje na cenu."""
    if not header:
        return False

    header_lower = str(header).lower()
    price_keywords = ['cena', 'cijena', 'price', 'cost', 'iznos', 'amount', 'vrednost', 'vrijednost']

    return any(keyword in header_lower for keyword in price_keywords)


def is_numeric_value(value):
    """Proverava da li je vrednost numerička."""
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        # Pokušaj da parsiraš string kao broj
        try:
            float(value.replace(',', '.'))
            return True
        except (ValueError, AttributeError):
            return False
    return False


def update_prices_in_sheet(sheet, coefficient, columns=None, dry_run=False):
    """
    Ažurira cene u sheet-u.

    Args:
        sheet: Openpyxl worksheet objekat
        coefficient: Koeficijent za množenje cena
        columns: Lista kolona za ažuriranje (npr. ['C', 'D']). Ako je None, automatski detektuje.
        dry_run: Ako je True, samo prikazuje šta bi se promenilo bez promena

    Returns:
        Broj ažuriranih ćelija
    """
    updated_count = 0

    # Ako kolone nisu specificirane, pokušaj automatsku detekciju
    if columns is None:
        columns = []
        if sheet.max_row > 0:
            # Pregledaj prvi red (header) za kolone sa cenama
            for col_idx in range(1, sheet.max_column + 1):
                header = sheet.cell(row=1, column=col_idx).value
                if is_price_column(header):
                    columns.append(get_column_letter(col_idx))

    # Ako i dalje nema kolona, koristi sve numeričke kolone
    if not columns:
        print(f"  ⚠️  Nisam pronašao kolone sa cenama u header-u, ažuriram sve numeričke vrednosti...")
        # Prolaziš kroz sve ćelije i ažuriraš numeričke
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                if is_numeric_value(cell.value):
                    old_value = cell.value
                    new_value = float(str(old_value).replace(',', '.')) * coefficient

                    if not dry_run:
                        cell.value = new_value

                    updated_count += 1
                    if dry_run:
                        print(f"    {cell.coordinate}: {old_value} -> {new_value:.2f}")
    else:
        # Ažuriraj specifične kolone
        print(f"  📊 Ažuriram kolone: {', '.join(columns)}")
        for col_letter in columns:
            for row_idx in range(2, sheet.max_row + 1):  # Preskačemo header (red 1)
                cell = sheet[f"{col_letter}{row_idx}"]

                if is_numeric_value(cell.value):
                    old_value = cell.value
                    new_value = float(str(old_value).replace(',', '.')) * coefficient

                    if not dry_run:
                        cell.value = new_value

                    updated_count += 1
                    if dry_run:
                        print(f"    {cell.coordinate}: {old_value} -> {new_value:.2f}")

    return updated_count


def update_excel_prices(file_path, coefficient, columns=None, sheets=None, dry_run=False, output_path=None):
    """
    Glavna funkcija za ažuriranje cena u Excel fajlu.

    Args:
        file_path: Putanja do Excel fajla
        coefficient: Koeficijent za množenje cena
        columns: Lista kolona za ažuriranje (npr. ['C', 'D']). Ako je None, automatski detektuje.
        sheets: Lista imena sheet-ova za ažuriranje. Ako je None, ažurira sve.
        dry_run: Ako je True, samo prikazuje šta bi se promenilo
        output_path: Putanja za čuvanje novog fajla. Ako je None, prepisuje originalni.
    """
    file_path = Path(file_path)

    if not file_path.exists():
        print(f"❌ Fajl '{file_path}' ne postoji!")
        sys.exit(1)

    print(f"📂 Učitavam fajl: {file_path}")

    try:
        workbook = load_workbook(filename=file_path)
    except Exception as e:
        print(f"❌ Greška pri učitavanju fajla: {e}")
        sys.exit(1)

    total_updated = 0

    # Odaberi koje sheet-ove ažurirati
    sheets_to_update = sheets if sheets else workbook.sheetnames

    print(f"\n🔄 Koeficijent: {coefficient}")
    print(f"📑 Sheet-ovi za ažuriranje: {len(sheets_to_update)}")

    if dry_run:
        print("\n⚠️  DRY RUN MODE - Nema stvarnih promena!\n")

    for sheet_name in sheets_to_update:
        if sheet_name not in workbook.sheetnames:
            print(f"⚠️  Sheet '{sheet_name}' ne postoji, preskačem...")
            continue

        sheet = workbook[sheet_name]
        print(f"\n📄 Sheet: {sheet_name}")

        updated = update_prices_in_sheet(sheet, coefficient, columns, dry_run)
        total_updated += updated

        print(f"  ✅ Ažurirano: {updated} ćelija")

    if not dry_run:
        # Sačuvaj fajl
        output = output_path if output_path else file_path
        print(f"\n💾 Čuvam izmene u: {output}")

        try:
            workbook.save(output)
            print(f"✅ Uspešno sačuvano!")
        except Exception as e:
            print(f"❌ Greška pri čuvanju: {e}")
            sys.exit(1)

    print(f"\n📊 Ukupno ažurirano ćelija: {total_updated}")

    if dry_run:
        print("\n💡 Da bi stvarno ažurirao cene, pokreni bez --dry-run opcije")


def main():
    parser = argparse.ArgumentParser(
        description='Automatski ažurira cene u Excel fajlovima sa više sheet-ova.',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Primeri korišćenja:

  # Povećaj sve cene za 10% (koeficijent 1.1)
  python update_prices.py cenovnik.xlsx 1.1

  # Smanji sve cene za 20% (koeficijent 0.8)
  python update_prices.py cenovnik.xlsx 0.8

  # Ažuriraj samo kolone C i D
  python update_prices.py cenovnik.xlsx 1.15 -c C D

  # Ažuriraj samo određene sheet-ove
  python update_prices.py cenovnik.xlsx 1.05 -s "Januar 2024" "Februar 2024"

  # Dry run - vidi šta bi se promenilo bez promena
  python update_prices.py cenovnik.xlsx 1.1 --dry-run

  # Sačuvaj u novi fajl
  python update_prices.py cenovnik.xlsx 1.1 -o cenovnik_updated.xlsx
        """
    )

    parser.add_argument('file', help='Putanja do Excel fajla')
    parser.add_argument('coefficient', type=float, help='Koeficijent za množenje cena (npr. 1.1 za +10%%, 0.9 za -10%%)')
    parser.add_argument('-c', '--columns', nargs='+', help='Kolone za ažuriranje (npr. C D E). Ako nije navedeno, automatski detektuje kolone sa cenama.')
    parser.add_argument('-s', '--sheets', nargs='+', help='Imena sheet-ova za ažuriranje. Ako nije navedeno, ažurira sve sheet-ove.')
    parser.add_argument('-o', '--output', help='Putanja za čuvanje novog fajla. Ako nije navedeno, prepisuje originalni fajl.')
    parser.add_argument('--dry-run', action='store_true', help='Prikaži šta bi se promenilo bez stvarnih promena')

    args = parser.parse_args()

    update_excel_prices(
        file_path=args.file,
        coefficient=args.coefficient,
        columns=args.columns,
        sheets=args.sheets,
        dry_run=args.dry_run,
        output_path=args.output
    )


if __name__ == '__main__':
    main()
